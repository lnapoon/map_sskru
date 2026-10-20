#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
สร้างเอกสาร Requirement (SRS) ของระบบ SSKRU Campus Map
เนื้อหาได้จากการวิเคราะห์ซอร์สโค้ดจริงในโปรเจกต์ (app.js, views.py, api.php,
migrate_schema.py, models.py, data/*.json)

ออกแบบให้เปิดได้บน Apple Pages (macOS):
  - ใช้ฟอนต์ Thonburi ที่มากับ macOS ทุกเครื่อง (ไม่มีหน้าต่างเตือนฟอนต์หาย)
  - ตารางใช้สไตล์ 'Table Grid' + สีพื้นหัวตารางเท่านั้น (OOXML ถูกต้องตาม schema)

Output: docs/SSKRU_Campus_Map_Requirements.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FONT = "Thonburi"
NAVY = RGBColor(0x0D, 0x2C, 0x5E)
BLUE = RGBColor(0x1A, 0x4F, 0xA0)


def _apply_font(run, size=12, bold=False, color=None):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # ขนาดฟอนต์สำหรับอักษรไทย (complex script)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)


def para(doc, text="", size=12, bold=False, align=None, color=None, before=0, after=4, indent=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.25
    if indent:
        pf.first_line_indent = Inches(0.4)
    if text:
        _apply_font(p.add_run(text), size, bold, color)
    return p


def h1(doc, text):
    para(doc, text, size=17, bold=True, color=NAVY, before=6, after=8)


def h2(doc, text):
    para(doc, text, size=14, bold=True, color=BLUE, before=10, after=4)


def h3(doc, text):
    para(doc, text, size=12.5, bold=True, before=6, after=2)


def bullets(doc, items, size=12):
    for it in items:
        p = para(doc, "", size=size, after=2)
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        _apply_font(p.add_run("•  " + it), size)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="1A4FA0"/>'))
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _apply_font(p.add_run(h), 11, True, RGBColor(0xFF, 0xFF, 0xFF))
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            lines = str(v).split("\n")
            p = cells[i].paragraphs[0]
            for j, line in enumerate(lines):
                if j:
                    p = cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                _apply_font(p.add_run(line), 10.5, bold=(i == 0))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    para(doc, "", after=4)
    return t


# ─────────────────────────────────────────────────────────────────────────────
def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Inches(11.69), Inches(8.27)  # A4
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)

    C = WD_ALIGN_PARAGRAPH.CENTER

    # ═══ ปก ═══
    logo = os.path.join(BASE_DIR, "images", "image1.png")
    if os.path.exists(logo):
        p = para(doc, align=C, before=20, after=16)
        p.add_run().add_picture(logo, width=Inches(1.2))
    para(doc, "เอกสารข้อกำหนดความต้องการของระบบ", 22, True, C, NAVY)
    para(doc, "Software Requirements Specification (SRS)", 14, True, C, BLUE, after=18)
    para(doc, "ระบบแผนผังดิจิทัลและระบบนำทางอัจฉริยะ", 18, True, C)
    para(doc, "มหาวิทยาลัยราชภัฏศรีสะเกษ (SSKRU Campus Map)", 18, True, C, after=40)
    para(doc, "จัดทำโดย", 14, True, C)
    para(doc, "นายมนต์พระกาฬ ทิมเกิด  รหัสนักศึกษา 6712732120", 14, False, C, after=20)
    para(doc, "เสนอ", 14, True, C)
    para(doc, "ผู้ช่วยศาสตราจารย์ ดร.กนิษฐา อินธิชิต", 14, False, C, after=40)
    para(doc, "รายวิชาระบบฐานข้อมูล (4123202)  ภาคเรียนที่ 2 ปีการศึกษา 2568", 12, False, C)
    para(doc, "สาขาวิทยาการคอมพิวเตอร์ คณะศิลปศาสตร์และวิทยาศาสตร์", 12, False, C)
    para(doc, "มหาวิทยาลัยราชภัฏศรีสะเกษ", 12, False, C)
    doc.add_page_break()

    # ═══ สารบัญ ═══
    para(doc, "สารบัญ", 18, True, C, NAVY, after=12)
    for line in [
        "1. บทนำและภาพรวมระบบ",
        "2. ผู้ใช้งานและสิทธิ์การเข้าถึง",
        "3. ความต้องการเชิงฟังก์ชัน (Functional Requirements)",
        "4. ความต้องการที่ไม่ใช่เชิงฟังก์ชัน (Non-Functional Requirements)",
        "5. สถาปัตยกรรมและเทคโนโลยีที่ใช้",
        "6. ความต้องการด้านข้อมูล (Data Requirements / Data Dictionary)",
        "7. ส่วนต่อประสานภายนอก (API Interfaces)",
        "8. กรณีการใช้งาน (Use Cases)",
        "9. ข้อจำกัด สมมติฐาน และข้อเสนอแนะจากการวิเคราะห์โค้ด",
        "10. ตารางตรวจสอบย้อนกลับ (Traceability Matrix)",
    ]:
        para(doc, line, 13, after=6)
    doc.add_page_break()

    # ═══ 1. บทนำ ═══
    h1(doc, "1. บทนำและภาพรวมระบบ")
    h2(doc, "1.1 วัตถุประสงค์ของเอกสาร")
    para(doc, "เอกสารนี้ระบุความต้องการของระบบ SSKRU Campus Map โดยวิเคราะห์จากซอร์สโค้ดที่พัฒนาจริงในโปรเจกต์ "
              "ทั้งส่วนหน้าเว็บ (index.php / index.html, assets/js/app.js, assets/css/styles.css) ส่วนแบ็กเอนด์ PHP (auth/*.php) "
              "ส่วนแบ็กเอนด์ Django (admin_panel, api, sskru_map) และส่วนฐานข้อมูล (database/migrate_schema.py, data/*.json)", indent=True)
    h2(doc, "1.2 ภาพรวมระบบ")
    para(doc, "ระบบเป็นเว็บแอปพลิเคชันแผนที่วิทยาเขตมหาวิทยาลัยราชภัฏศรีสะเกษ แสดงแผนผัง 3D ของมหาวิทยาลัย (images/Map.png) "
              "บนไลบรารี Leaflet พร้อมหมุดอาคาร 40 แห่ง ค้นหาอาคาร ระบุตำแหน่งผู้ใช้ด้วย GPS คำนวณเส้นทางด้วยอัลกอริทึม Dijkstra "
              "และส่งต่อการนำทางไปยัง Google Maps มีระบบสมาชิกหลายบทบาท (บุคคลทั่วไป นักศึกษา บุคลากร ผู้ดูแลระบบ) "
              "และแผงผู้ดูแลระบบสำหรับจัดการข้อมูลอาคารและดูสถิติการใช้งาน", indent=True)
    h2(doc, "1.3 ขอบเขตของระบบ")
    bullets(doc, [
        "ข้อมูลอาคาร 40 แห่ง แบ่ง 5 หมวด: สิ่งอำนวยความสะดวก 18, อาคารเรียน/คณะ 10, สถานที่อื่นๆ 7, สำนักงาน 3, หอสมุด 2",
        "ทะเบียนนักศึกษาตัวอย่าง (students_roster.json) 30 รายการ บัญชีนักศึกษาที่ลงทะเบียนแล้ว 8 บัญชี บัญชีบุคลากร 2 บัญชี",
        "รองรับการติดตั้ง 2 แบบ: XAMPP (PHP + MySQL) และ Django (Python + SQLite) และ deploy บน Vercel",
        "ไม่ครอบคลุม: ระบบลงทะเบียนเรียน ตารางสอน หรือการเชื่อมต่อระบบทะเบียนจริงของมหาวิทยาลัย",
    ])
    h2(doc, "1.4 นิยามศัพท์")
    table(doc, ["คำศัพท์", "ความหมาย"], [
        ["Guest Mode", "โหมดผู้เยี่ยมชม/บุคคลทั่วไป เข้าดูแผนที่ได้โดยไม่ต้องมีบัญชี และไม่มีข้อมูลนักศึกษา"],
        ["Roster", "ทะเบียนรายชื่อนักศึกษาที่ใช้ตรวจสอบสิทธิ์ก่อนลงทะเบียน (รหัสนักศึกษา + เลขบัตรประชาชน)"],
        ["OTP", "รหัสผ่านใช้ครั้งเดียว 6 หลัก ส่งทางอีเมล ใช้ยืนยันการรีเซ็ตรหัสผ่าน (อายุ 15 นาที)"],
        ["PDPA", "พ.ร.บ.คุ้มครองข้อมูลส่วนบุคคล พ.ศ. 2562"],
        ["Dijkstra", "อัลกอริทึมหาเส้นทางสั้นที่สุดบนกราฟจุดเชื่อมต่อภายในวิทยาเขต"],
        ["PWA", "Progressive Web App ติดตั้งลงหน้าจอมือถือได้ (manifest.json + service worker)"],
    ], [1.6, 4.6])
    doc.add_page_break()

    # ═══ 2. ผู้ใช้งาน ═══
    h1(doc, "2. ผู้ใช้งานและสิทธิ์การเข้าถึง")
    table(doc, ["บทบาท", "การเข้าสู่ระบบ", "สิทธิ์ที่ได้รับ", "ข้อจำกัด"], [
        ["บุคคลทั่วไป /\nผู้เยี่ยมชม", "กดปุ่ม 'บุคคลทั่วไป เข้าดูแผนที่ได้เลย' ที่หน้าล็อกอิน (ไม่ต้องใช้รหัสผ่าน)",
         "ดูแผนที่ ค้นหาอาคาร ดูรายละเอียด ระบุ GPS นำทาง แชร์ตำแหน่ง",
         "ไม่มีข้อมูลนักศึกษา/ข้อมูลส่วนตัว ไม่มีฟอร์มแก้ไขโปรไฟล์ ไม่มีปุ่มปลดล็อกข้อมูล"],
        ["นักศึกษา", "รหัสนักศึกษา + รหัสผ่าน (ลงทะเบียนครั้งแรกด้วยรหัสนักศึกษา + เลขบัตรประชาชน)",
         "ทุกอย่างของบุคคลทั่วไป + โปรไฟล์ (Bio, ชื่อเล่น) + ดูข้อมูลส่วนตัวหลังยืนยันรหัสผ่าน",
         "ต้องมีรายชื่อใน Roster ก่อนลงทะเบียน"],
        ["บุคลากร /\nอาจารย์", "Username หรือ Email + รหัสผ่าน",
         "ดูแผนที่ ค้นหา นำทาง และโปรไฟล์บุคลากร",
         "ต้องรอผู้ดูแลระบบอนุมัติบัญชี (is_approved) ก่อนเข้าใช้งาน"],
        ["ผู้ดูแลระบบ", "บัญชีผู้ดูแลระบบ (กำหนดใน .env: ADMIN_USERNAME / ADMIN_PASSWORD)",
         "เพิ่ม/แก้ไข/ลบอาคาร ลากย้ายหมุด จัดการนักศึกษาและบุคลากร ดูสถิติ Dashboard",
         "เซสชันหมดอายุภายใน 8 ชั่วโมง"],
    ], [1.1, 1.7, 1.9, 1.6])
    doc.add_page_break()

    # ═══ 3. Functional ═══
    h1(doc, "3. ความต้องการเชิงฟังก์ชัน (Functional Requirements)")
    para(doc, "ระดับความสำคัญ: สูง = ต้องมี, กลาง = ควรมี, ต่ำ = มีได้ถ้าเหมาะสม  ·  อ้างอิง = ฟังก์ชัน/ไฟล์ที่พัฒนาความต้องการนั้นจริง", 11)
    W = [0.9, 1.6, 0.6, 3.1]
    HD = ["รหัส", "ความต้องการ", "ระดับ", "รายละเอียด / อ้างอิงโค้ด"]

    h2(doc, "3.1 การแสดงผลแผนที่")
    table(doc, HD, [
        ["FR-01", "แสดงแผนผัง 3D ของวิทยาเขต", "สูง", "แสดงภาพ Map.png บน Leaflet (initMap) ซูม/เลื่อน/รีเซ็ตมุมมองได้"],
        ["FR-02", "แสดงหมุดอาคารตามหมวดหมู่", "สูง", "renderMarkers แสดงหมุดสีตามหมวด 5 หมวด พร้อมเลขอาคาร"],
        ["FR-03", "แถบการ์ดอาคาร (Carousel)", "กลาง", "renderBuildingCarousel แตะการ์ดแล้วแผนที่เลื่อนไปยังอาคาร"],
        ["FR-04", "ทัวร์ 3D อัตโนมัติ", "ต่ำ", "toggle3dTour / stop3dTour พาชมอาคารต่อเนื่อง"],
        ["FR-05", "สถานะเปิด-ปิดตามเวลาจริง", "กลาง", "getBuildingStatus / updateRealTimeStatus คำนวณจากวันและเวลาปัจจุบัน"],
    ], W)

    h2(doc, "3.2 การค้นหาและแสดงรายละเอียดอาคาร")
    table(doc, HD, [
        ["FR-06", "ค้นหาอาคาร", "สูง", "handleSearchInput / handleMobileSearchInput ค้นจากชื่อไทย อังกฤษ รหัส และแท็ก"],
        ["FR-07", "กรองตามหมวดหมู่", "สูง", "filterByCategory: ทั้งหมด, อาคารเรียน, สำนักงาน, สิ่งอำนวยความสะดวก, หอสมุด, อื่นๆ"],
        ["FR-08", "แผงรายละเอียดอาคาร", "สูง", "selectBuilding แสดงชื่อ รหัส คำอธิบาย เบอร์โทร สถานะ และปุ่มนำทาง"],
        ["FR-09", "แชร์ตำแหน่งอาคาร", "กลาง", "shareBuildingLocation ใช้ Web Share API / คัดลอกลิงก์ Google Maps"],
    ], W)

    h2(doc, "3.3 การระบุตำแหน่งและนำทาง")
    table(doc, HD, [
        ["FR-10", "ระบุตำแหน่งผู้ใช้ด้วย GPS", "สูง", "trackUserLocation ใช้ Geolocation API แสดงจุดตำแหน่งบนแผนที่"],
        ["FR-11", "แปลงพิกัด GPS เป็นพิกัดแผนที่", "สูง", "gpsToMap แปลง Lat/Lng จริงเป็นพิกัด X/Y บนภาพแผนผัง"],
        ["FR-12", "คำนวณเส้นทางภายในวิทยาเขต", "สูง", "buildNetworkGraph + solveDijkstra หาเส้นทางสั้นที่สุด calculateHaversineDistance คำนวณระยะทาง"],
        ["FR-13", "เลือกวิธีเดินทางและประมาณเวลา", "กลาง", "เดินเท้า 1.2 ม./วิ, จักรยาน 3.5, มอเตอร์ไซค์ 7.5, รถยนต์ 6.0 (TRAVEL_MODES)"],
        ["FR-14", "จำลองการเดินทางบนเส้นทาง", "ต่ำ", "startGrabRiderMotion แสดงไอคอนเคลื่อนที่ตามเส้นทาง และแบนเนอร์เมื่อถึงที่หมาย"],
        ["FR-15", "นำทาง Turn-by-Turn / Google Maps", "สูง", "startActiveTurnByTurnNav และลิงก์ไป Google Maps ตามโหมดการเดินทาง"],
    ], W)

    h2(doc, "3.4 การยืนยันตัวตนและบัญชีผู้ใช้")
    table(doc, HD, [
        ["FR-16", "หน้าล็อกอินรวม 3 แท็บ", "สูง", "auth/login.php, student_login.html: แท็บนักศึกษา / บุคลากร / บุคคลทั่วไป"],
        ["FR-17", "เข้าใช้งานแบบบุคคลทั่วไป", "สูง", "action guest_login (PHP) / guest_login_api (Django) ตั้ง session role = guest แล้วเข้าแผนที่ทันที"],
        ["FR-18", "ลงทะเบียนนักศึกษา", "สูง", "student_verify ตรวจรหัสนักศึกษา + เลขบัตรประชาชนกับ Roster แล้ว student_register ตั้งรหัสผ่าน"],
        ["FR-19", "เข้าสู่ระบบนักศึกษา", "สูง", "student_login รองรับรหัสแบบมีขีด/ไม่มีขีด"],
        ["FR-20", "สมัครและเข้าสู่ระบบบุคลากร", "สูง", "staff_register / staff_login เข้าได้หลังผู้ดูแลระบบอนุมัติ"],
        ["FR-21", "รีเซ็ตรหัสผ่านด้วย OTP", "สูง", "request_reset ส่ง OTP 6 หลักทางอีเมล (SMTP) อายุ 15 นาที → verify_reset_token → confirm_new_password"],
        ["FR-22", "จำกัดการพยายามเข้าสู่ระบบ", "สูง", "check_rate_limit: ผิด 5 ครั้งล็อก 10 นาที (ล็อกอิน/OTP) และ 5 นาที (ยืนยันรหัสผ่านดูข้อมูล)"],
        ["FR-23", "ออกจากระบบ", "สูง", "handleLogout / action logout ล้าง session (โหมดผู้เยี่ยมชมแสดงเป็น 'กลับไปหน้าเข้าสู่ระบบ')"],
    ], W)

    h2(doc, "3.5 โปรไฟล์และความเป็นส่วนตัว")
    table(doc, HD, [
        ["FR-24", "แสดงข้อมูลผู้ใช้ในเมนูด้านข้าง", "กลาง", "updateDrawerUserInfo แสดงชื่อและบทบาท (ผู้เยี่ยมชมแสดง 'ผู้เยี่ยมชมมหาวิทยาลัย')"],
        ["FR-25", "แผงโปรไฟล์", "กลาง", "populateProfilePanel แก้ไข Bio (120 ตัวอักษร) และชื่อเล่น (30 ตัวอักษร)"],
        ["FR-26", "พรางข้อมูลส่วนบุคคล", "สูง", "updateProfilePrivacyDisplay ซ่อนชื่อจริง รหัส อีเมล เป็น •••• ต้องยืนยันรหัสผ่านก่อนแสดง (verify_password)"],
        ["FR-27", "แยกข้อมูลผู้เยี่ยมชม", "สูง", "Guest Mode ซ่อนฟอร์มแก้ไขโปรไฟล์และปุ่มปลดล็อก แสดงข้อมูลเป็น 'บุคคลทั่วไป' ไม่มีรหัสนักศึกษา/อีเมล"],
        ["FR-28", "ขอความยินยอม PDPA", "สูง", "initPDPAConsent / openPDPAModal ใช้นโยบายจาก data/pdpa_policy.json"],
    ], W)

    h2(doc, "3.6 ผู้ดูแลระบบ")
    table(doc, HD, [
        ["FR-29", "เพิ่ม/แก้ไข/ลบอาคาร", "สูง", "openAddBuildingForm / openEditBuildingForm / handleBuildingDelete ผ่าน admin_buildings_api"],
        ["FR-30", "ลากย้ายหมุดอาคาร", "สูง", "enterAdminMode เปิดการลากหมุด บันทึกพิกัดใหม่ทันที"],
        ["FR-31", "ซิงก์ข้อมูลหลายแหล่ง", "กลาง", "write_buildings บันทึก buildings.json แล้วซิงก์ MySQL / MongoDB เบื้องหลัง (thread)"],
        ["FR-32", "จัดการนักศึกษาและบุคลากร", "สูง", "admin_students_api / admin_staff_api เพิ่ม ลบ อนุมัติบัญชีบุคลากร"],
        ["FR-33", "Dashboard สถิติ", "กลาง", "admin_analytics_api สรุปผู้เข้าชม อุปกรณ์ เหตุการณ์ (ค้นหา เลือกอาคาร นำทาง GPS แชร์)"],
        ["FR-34", "ส่งออกข้อมูลอาคาร", "ต่ำ", "openExportModal ส่งออกข้อมูลอาคารเป็น JSON"],
        ["FR-35", "บันทึกการใช้งาน", "กลาง", "trackEvent / track_event และ log_user_activity บันทึก IP อุปกรณ์ เบราว์เซอร์"],
    ], W)

    h2(doc, "3.7 ส่วนติดต่อผู้ใช้")
    table(doc, HD, [
        ["FR-36", "โหมดมืด/สว่าง", "กลาง", "toggleTheme จำค่าไว้ใน localStorage (sskru_theme)"],
        ["FR-37", "รองรับมือถือ", "สูง", "Responsive layout, setupSwipeGesture, เมนูด้านข้าง (side drawer)"],
        ["FR-38", "ติดตั้งเป็นแอป (PWA)", "ต่ำ", "assets/manifest.json + sw.js"],
        ["FR-39", "กล่องข้อความแจ้งเตือนแบบกำหนดเอง", "ต่ำ", "dialog.js: showSskruAlert / showSskruConfirm แทน alert() ของเบราว์เซอร์"],
    ], W)
    doc.add_page_break()

    # ═══ 4. NFR ═══
    h1(doc, "4. ความต้องการที่ไม่ใช่เชิงฟังก์ชัน (Non-Functional Requirements)")
    table(doc, ["รหัส", "ด้าน", "ความต้องการ / เกณฑ์"], [
        ["NFR-01", "ประสิทธิภาพ", "แสดงหมุดอาคารจากแคช localStorage ทันทีที่เปิดหน้า แล้วค่อยซิงก์จากเซิร์ฟเวอร์เบื้องหลัง (timeout 2 วินาที)"],
        ["NFR-02", "ประสิทธิภาพ", "การค้นหาต้องแสดงผลทันทีขณะพิมพ์ โดยไม่ต้องเรียกเซิร์ฟเวอร์"],
        ["NFR-03", "ความพร้อมใช้งาน", "ถ้า API ไม่ตอบสนอง ต้องอ่านข้อมูลสำรองตามลำดับ: API → data/buildings.json → ข้อมูลในโค้ด (BUILDINGS)"],
        ["NFR-04", "ความปลอดภัย", "รหัสผ่านต้องเก็บแบบแฮช (ปัจจุบัน SHA-256) และไม่แสดงในหน้าเว็บ"],
        ["NFR-05", "ความปลอดภัย", "จำกัดการเดารหัสผ่าน/OTP สูงสุด 5 ครั้ง แล้วล็อกชั่วคราว"],
        ["NFR-06", "ความปลอดภัย", "ห้ามเข้าถึงไฟล์ .json .sqlite3 .env .log และโฟลเดอร์ data/ โดยตรง (.htaccess)"],
        ["NFR-07", "ความปลอดภัย", "โทเค็นผู้ดูแลระบบหมดอายุ 8 ชั่วโมง และ OTP หมดอายุ 15 นาที ใช้ได้ครั้งเดียว"],
        ["NFR-08", "ความเป็นส่วนตัว", "ผู้เยี่ยมชมต้องไม่เห็นหรือถูกผูกกับข้อมูลนักศึกษาใดๆ ข้อมูลส่วนตัวนักศึกษาพรางไว้เป็นค่าเริ่มต้น (PDPA)"],
        ["NFR-09", "การใช้งาน", "ภาษาไทยเป็นหลัก ใช้ฟอนต์ Sarabun/Outfit ผู้ใช้ทั่วไปเข้าแผนที่ได้ภายใน 1 คลิกจากหน้าล็อกอิน"],
        ["NFR-10", "ความเข้ากันได้", "ทำงานบน Chrome, Safari, Edge, Firefox ทั้งเดสก์ท็อป iOS และ Android"],
        ["NFR-11", "การติดตั้ง", "ติดตั้งได้ทั้ง XAMPP (PHP 7.4+ / MySQL) และ Django 4.2+ (Python 3) และ Vercel"],
        ["NFR-12", "การบำรุงรักษา", "แยกส่วน Frontend / API / ฐานข้อมูล ตั้งค่าความลับผ่านไฟล์ .env"],
    ], [0.9, 1.2, 4.1])
    doc.add_page_break()

    # ═══ 5. สถาปัตยกรรม ═══
    h1(doc, "5. สถาปัตยกรรมและเทคโนโลยีที่ใช้")
    table(doc, ["ส่วน", "เทคโนโลยี", "ไฟล์หลัก"], [
        ["Frontend", "HTML5, CSS3 (Vanilla), JavaScript ES6, Leaflet 1.9.4, Font Awesome 6.4, Google Fonts", "index.php, index.html, assets/js/app.js, assets/css/styles.css, assets/js/dialog.js"],
        ["Backend (PHP)", "PHP + PDO + MySQL (XAMPP / phpMyAdmin)", "auth/api.php, auth/config.php, auth/login.php, auth/register_*.php, auth/reset_password_*.php"],
        ["Backend (Python)", "Django 4.2, python-dotenv, Pillow, dj-database-url", "sskru_map/settings.py, admin_panel/views.py, api/views.py"],
        ["ฐานข้อมูล", "MySQL (8 ตาราง), SQLite (Django models), ไฟล์ JSON, MongoDB (ซิงก์สำรอง)", "database/migrate_schema.py, admin_panel/models.py, data/*.json"],
        ["อีเมล", "SMTP (Gmail) สำหรับส่ง OTP", "EMAIL_* ใน .env / settings.py"],
        ["Deploy", "Vercel (@vercel/python), Apache .htaccess", "vercel.json, deploy/build_vercel.sh"],
    ], [1.2, 2.5, 2.5])
    h2(doc, "5.1 ลำดับการทำงานเมื่อเปิดระบบ")
    bullets(doc, [
        "ผู้ใช้เปิดหน้าแรก → ถ้ายังไม่มี session จะถูกส่งไปหน้าล็อกอิน (/login/ หรือ auth/login.php)",
        "เลือกบทบาท: นักศึกษา / บุคลากร / บุคคลทั่วไป → เซิร์ฟเวอร์ตั้งค่า session ตามบทบาท",
        "หน้าแผนที่เรียก auth_status เพื่อรู้บทบาท แล้วปรับเมนูและโปรไฟล์ให้ตรงสิทธิ์",
        "โหลดข้อมูลอาคารจากแคชทันที แล้วซิงก์จาก API / buildings.json เบื้องหลัง",
    ])
    doc.add_page_break()

    # ═══ 6. Data ═══
    h1(doc, "6. ความต้องการด้านข้อมูล (Data Dictionary)")
    para(doc, "โครงสร้างฐานข้อมูล MySQL ตาม database/migrate_schema.py มี 8 ตาราง (Engine InnoDB, utf8mb4)", indent=True)
    erd = os.path.join(BASE_DIR, "images", "er_diagram_hd.png")
    if os.path.exists(erd):
        p = para(doc, align=C, after=2)
        p.add_run().add_picture(erd, width=Inches(6.0))
        para(doc, "รูปที่ 1 แผนภาพความสัมพันธ์ของข้อมูล (E-R Diagram)", 10.5, True, C, after=8)

    DW = [1.4, 1.5, 1.1, 2.2]
    DH = ["ฟิลด์", "ชนิดข้อมูล", "คีย์/ข้อกำหนด", "คำอธิบาย"]
    h3(doc, "6.1 categories — หมวดหมู่อาคาร")
    table(doc, DH, [
        ["category_key", "VARCHAR(50)", "PK", "รหัสหมวด: academic, office, facility, library, other"],
        ["category_name_th", "VARCHAR(100)", "NOT NULL", "ชื่อหมวดภาษาไทย"],
        ["category_name_en", "VARCHAR(100)", "NOT NULL", "ชื่อหมวดภาษาอังกฤษ"],
        ["color_hex", "VARCHAR(20)", "NOT NULL", "สีหมุดบนแผนที่"],
        ["icon_class", "VARCHAR(50)", "DEFAULT fa-building", "ไอคอน Font Awesome"],
        ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันที่สร้าง"],
    ], DW)
    h3(doc, "6.2 faculties — คณะ/หน่วยงาน")
    table(doc, DH, [
        ["faculty_id", "INT", "PK, AUTO_INCREMENT", "รหัสคณะ"],
        ["faculty_name", "VARCHAR(150)", "UNIQUE, NOT NULL", "ชื่อคณะ/หน่วยงาน (8 รายการเริ่มต้น)"],
        ["faculty_code", "VARCHAR(50)", "", "อักษรย่อ เช่น FLAS, EDU"],
        ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันที่สร้าง"],
    ], DW)
    h3(doc, "6.3 buildings — อาคารและสถานที่")
    table(doc, DH, [
        ["id", "INT", "PK", "หมายเลขอาคาร"],
        ["code", "VARCHAR(20)", "", "รหัสย่ออาคาร"],
        ["name / name_en", "VARCHAR(255)", "name NOT NULL", "ชื่ออาคารไทย / อังกฤษ"],
        ["category_key", "VARCHAR(50)", "FK → categories", "หมวดหมู่อาคาร"],
        ["coords_x / coords_y", "INT", "NOT NULL", "พิกัดบนภาพแผนผัง"],
        ["real_lat / real_lng", "DOUBLE", "", "พิกัด GPS จริง"],
        ["description", "TEXT", "", "รายละเอียดอาคาร"],
        ["phone / hours", "VARCHAR", "", "เบอร์โทร / เวลาทำการ"],
        ["floors_count", "INT", "DEFAULT 1", "จำนวนชั้น"],
        ["json_data", "LONGTEXT", "", "ข้อมูลเต็มรูปแบบ JSON (tags ฯลฯ)"],
        ["created_at / updated_at", "TIMESTAMP", "", "วันที่สร้าง / แก้ไขล่าสุด"],
    ], DW)
    h3(doc, "6.4 students_roster — ทะเบียนนักศึกษา")
    table(doc, DH, [
        ["student_id", "VARCHAR(20)", "PK", "รหัสนักศึกษา"],
        ["citizen_id", "VARCHAR(20)", "", "เลขบัตรประชาชน (ใช้ยืนยันตอนลงทะเบียน)"],
        ["name / name_en", "VARCHAR(150)", "name NOT NULL", "ชื่อ-นามสกุล"],
        ["faculty_id", "INT", "FK → faculties", "คณะที่สังกัด"],
        ["major", "VARCHAR(150)", "", "สาขาวิชา"],
        ["status", "ENUM", "studying/graduated/resigned", "สถานะนักศึกษา"],
        ["created_at", "TIMESTAMP", "DEFAULT NOW", "วันที่สร้าง"],
    ], DW)
    h3(doc, "6.5 users — บัญชีผู้ใช้ (ผู้ดูแลระบบ / บุคลากร / นักศึกษา)")
    table(doc, DH, [
        ["user_id", "INT", "PK, AUTO_INCREMENT", "รหัสบัญชี"],
        ["username", "VARCHAR(100)", "UNIQUE, NOT NULL", "ชื่อผู้ใช้"],
        ["email", "VARCHAR(150)", "UNIQUE", "อีเมล"],
        ["password_hash", "VARCHAR(255)", "NOT NULL", "รหัสผ่านแบบแฮช SHA-256"],
        ["role", "ENUM", "admin/staff/student", "บทบาท"],
        ["student_id", "VARCHAR(20)", "FK → students_roster", "ผูกกับทะเบียนนักศึกษา (กรณีนักศึกษา)"],
        ["display_name / bio", "VARCHAR / TEXT", "", "ชื่อที่แสดง / คำแนะนำตัว"],
        ["is_active", "BOOLEAN", "DEFAULT TRUE", "สถานะการใช้งาน"],
        ["created_at / updated_at", "TIMESTAMP", "", "วันที่สร้าง / แก้ไข"],
    ], DW)
    h3(doc, "6.6 user_activity_logs — ประวัติการเข้าใช้งาน")
    table(doc, DH, [
        ["log_id", "BIGINT", "PK, AUTO_INCREMENT", "รหัสบันทึก"],
        ["user_id", "INT", "FK → users", "ผู้ใช้ (NULL ได้ เช่น ผู้เยี่ยมชม)"],
        ["identifier / name", "VARCHAR", "NOT NULL", "รหัสหรือชื่อผู้ใช้ (ผู้เยี่ยมชม = GUEST)"],
        ["role / role_th", "VARCHAR", "", "บทบาท (guest, student, staff, admin)"],
        ["email", "VARCHAR(150)", "", "อีเมล"],
        ["ip_address / device / os_name / browser", "VARCHAR(50)", "", "ข้อมูลเครือข่ายและอุปกรณ์"],
        ["action", "VARCHAR(100)", "", "กิจกรรม เช่น เข้าสู่ระบบ"],
        ["created_at", "DATETIME", "DEFAULT NOW", "วันเวลา"],
    ], DW)
    h3(doc, "6.7 password_resets — คำขอรีเซ็ตรหัสผ่าน")
    table(doc, DH, [
        ["reset_id", "BIGINT", "PK, AUTO_INCREMENT", "รหัสคำขอ"],
        ["student_id / username", "VARCHAR", "", "ผู้ขอ (นักศึกษาหรือบุคลากร)"],
        ["email", "VARCHAR(150)", "NOT NULL", "อีเมลที่ส่ง OTP"],
        ["token", "VARCHAR(64)", "UNIQUE", "โทเค็นอ้างอิงคำขอ"],
        ["otp", "VARCHAR(10)", "NOT NULL", "รหัส OTP 6 หลัก"],
        ["expires_at", "DATETIME", "NOT NULL", "เวลาหมดอายุ (15 นาที)"],
        ["is_used", "BOOLEAN", "DEFAULT FALSE", "ใช้งานแล้วหรือยัง"],
        ["created_at", "DATETIME", "DEFAULT NOW", "วันเวลาที่ขอ"],
    ], DW)
    h3(doc, "6.8 admin_sessions — เซสชันผู้ดูแลระบบ")
    table(doc, DH, [
        ["session_id", "BIGINT", "PK, AUTO_INCREMENT", "รหัสเซสชัน"],
        ["token", "VARCHAR(128)", "UNIQUE", "โทเค็นยืนยันตัวตน"],
        ["user_id", "INT", "FK → users (CASCADE)", "ผู้ดูแลระบบเจ้าของเซสชัน"],
        ["is_active", "BOOLEAN", "DEFAULT TRUE", "สถานะเซสชัน"],
        ["expires_at", "DATETIME", "NOT NULL", "เวลาหมดอายุ (8 ชั่วโมง)"],
        ["created_at", "DATETIME", "DEFAULT NOW", "วันเวลาที่สร้าง"],
    ], DW)
    h3(doc, "6.9 ความสัมพันธ์ระหว่างตาราง")
    bullets(doc, [
        "categories 1 : N buildings  (category_key)",
        "faculties 1 : N students_roster  (faculty_id)",
        "students_roster 1 : 0..1 users  (student_id)",
        "users 1 : N user_activity_logs  (user_id)",
        "users 1 : N admin_sessions  (user_id)",
        "users / students_roster 1 : N password_resets  (username / student_id)",
    ])
    h3(doc, "6.10 ตารางเสริมฝั่ง Django (SQLite)")
    para(doc, "admin_panel/models.py มีโมเดล VisitorLog (สถิติผู้เข้าชม), UserEvent (เหตุการณ์ page_view, building_select, search, navigate, gps_track, share), "
              "AdminSession, Student, StaffUser (มี is_approved), PasswordResetToken, Building และ UserActivityLog ใช้งานเมื่อติดตั้งแบบ Django", indent=True)
    doc.add_page_break()

    # ═══ 7. API ═══
    h1(doc, "7. ส่วนต่อประสานภายนอก (API Interfaces)")
    table(doc, ["ฟังก์ชัน", "PHP (auth/api.php?action=)", "Django"], [
        ["สถานะการเข้าสู่ระบบ", "auth_status", "GET /admin/api/auth/status/"],
        ["เข้าใช้แบบบุคคลทั่วไป", "guest_login", "POST /admin/api/auth/guest_login/"],
        ["ล็อกอินนักศึกษา", "student_login", "POST /admin/api/auth/student_login/"],
        ["ตรวจสอบ / ลงทะเบียนนักศึกษา", "student_verify, student_register", "/admin/api/auth/student_verify/, student_register/"],
        ["สมัคร / ล็อกอินบุคลากร", "staff_register, staff_login", "/admin/api/auth/staff_register/, staff_login/"],
        ["รีเซ็ตรหัสผ่าน (3 ขั้น)", "*_request_reset, *_verify_reset_token, *_confirm_new_password", "/admin/api/auth/<role>_request_reset/ ฯลฯ"],
        ["ยืนยันรหัสผ่านเพื่อดูข้อมูล", "verify_password", "POST /admin/api/auth/verify_password/"],
        ["ออกจากระบบ", "logout", "GET /logout/"],
        ["ข้อมูลอาคาร (อ่าน)", "GET auth/api.php", "GET /api/buildings"],
        ["จัดการอาคาร (ผู้ดูแล)", "POST/PUT/DELETE auth/api.php?id=", "/admin/api/buildings/, /admin/api/buildings/<id>/"],
        ["สถิติ / นักศึกษา / บุคลากร", "-", "/admin/api/analytics/, students/, staff/"],
        ["บันทึกเหตุการณ์", "-", "POST /api/track"],
        ["ตรวจสุขภาพระบบ", "-", "GET /api/health"],
    ], [1.8, 2.3, 2.1])
    doc.add_page_break()

    # ═══ 8. Use cases ═══
    h1(doc, "8. กรณีการใช้งาน (Use Cases)")
    def uc(code, name, actor, pre, steps, post):
        h3(doc, f"{code}: {name}")
        table(doc, ["หัวข้อ", "รายละเอียด"], [
            ["ผู้ใช้งาน", actor], ["เงื่อนไขก่อน", pre],
            ["ขั้นตอนหลัก", "\n".join(f"{i}. {s}" for i, s in enumerate(steps, 1))],
            ["ผลลัพธ์", post],
        ], [1.3, 4.9])
    uc("UC-01", "บุคคลทั่วไปเข้าดูแผนที่", "ผู้ปกครอง / แขก / ผู้ใหญ่ที่เข้ามาเยี่ยมชม", "ไม่มีบัญชี",
       ["เปิดเว็บไซต์ ระบบแสดงหน้าล็อกอิน", "กด 'บุคคลทั่วไป เข้าดูแผนที่ได้เลย' หรือแท็บ 'บุคคลทั่วไป'",
        "ระบบตั้ง session role = guest และบันทึกกิจกรรม (GUEST)", "ระบบเปิดหน้าแผนที่ เมนูแสดง 'ผู้เยี่ยมชมมหาวิทยาลัย'"],
       "ใช้แผนที่ ค้นหา และนำทางได้ครบ โดยไม่มีข้อมูลนักศึกษาใดๆ")
    uc("UC-02", "ค้นหาอาคารและนำทาง", "ทุกบทบาท", "เข้าหน้าแผนที่แล้ว",
       ["พิมพ์ชื่ออาคารหรือเลือกหมวดหมู่", "เลือกอาคารจากผลลัพธ์ แผนที่เลื่อนไปยังอาคาร",
        "กดปุ่ม GPS อนุญาตตำแหน่ง", "เลือกวิธีเดินทาง ระบบคำนวณเส้นทาง ระยะทาง และเวลา", "กดนำทาง หรือเปิดใน Google Maps"],
       "ผู้ใช้เห็นเส้นทางและไปถึงอาคารเป้าหมาย")
    uc("UC-03", "นักศึกษาลงทะเบียนและเข้าสู่ระบบ", "นักศึกษา", "มีรายชื่อใน Roster",
       ["เลือก 'ลงทะเบียนครั้งแรก' กรอกรหัสนักศึกษาและเลขบัตรประชาชน", "ระบบตรวจสอบกับ Roster",
        "ตั้งรหัสผ่าน", "เข้าสู่ระบบด้วยรหัสนักศึกษาและรหัสผ่าน"],
       "เข้าแผนที่ในบทบาทนักศึกษา ข้อมูลส่วนตัวถูกพรางไว้")
    uc("UC-04", "รีเซ็ตรหัสผ่านด้วย OTP", "นักศึกษา / บุคลากร", "มีบัญชีและอีเมล",
       ["กด 'ลืมรหัสผ่าน?' กรอกรหัส/อีเมล", "ระบบส่ง OTP 6 หลักทางอีเมล", "กรอก OTP ภายใน 15 นาที (ผิดได้ไม่เกิน 5 ครั้ง)", "ตั้งรหัสผ่านใหม่"],
       "รหัสผ่านถูกเปลี่ยน OTP ถูกทำเครื่องหมายว่าใช้แล้ว")
    uc("UC-05", "บุคลากรสมัครและรออนุมัติ", "บุคลากร, ผู้ดูแลระบบ", "-",
       ["บุคลากรกรอก Username, Email, รหัสผ่าน", "ระบบบันทึกบัญชีสถานะรออนุมัติ", "ผู้ดูแลระบบอนุมัติใน Dashboard", "บุคลากรเข้าสู่ระบบได้"],
       "บัญชีบุคลากรใช้งานได้")
    uc("UC-06", "ผู้ดูแลระบบจัดการอาคาร", "ผู้ดูแลระบบ", "เข้าสู่ระบบผู้ดูแลแล้ว",
       ["เปิดโหมดผู้ดูแล", "ลากหมุดเพื่อปรับพิกัด หรือเปิดฟอร์มเพิ่ม/แก้ไข/ลบ", "บันทึก", "ระบบเขียน buildings.json และซิงก์ฐานข้อมูลอื่น"],
       "ข้อมูลอาคารอัปเดตสำหรับผู้ใช้ทุกคน")
    doc.add_page_break()

    # ═══ 9. Constraints & findings ═══
    h1(doc, "9. ข้อจำกัด สมมติฐาน และข้อเสนอแนะจากการวิเคราะห์โค้ด")
    h2(doc, "9.1 สมมติฐานและข้อจำกัด")
    bullets(doc, [
        "การระบุตำแหน่งต้องได้รับอนุญาตจากผู้ใช้ และความแม่นยำขึ้นกับ GPS ของอุปกรณ์",
        "เบราว์เซอร์ส่วนใหญ่อนุญาต Geolocation เฉพาะ HTTPS หรือ localhost",
        "การส่ง OTP ต้องตั้งค่า SMTP ใน .env ให้ถูกต้อง",
        "ข้อมูลหลักของอาคารอยู่ใน data/buildings.json การซิงก์ MySQL/MongoDB เป็นแบบเบื้องหลัง ถ้าเชื่อมต่อไม่ได้จะข้ามไป",
    ])
    h2(doc, "9.2 ประเด็นที่ควรปรับปรุง (พบจากการตรวจโค้ด)")
    table(doc, ["ประเด็น", "ตำแหน่งที่พบ", "ข้อเสนอแนะ"], [
        ["มีรหัสผ่านฐานข้อมูล MongoDB และรหัสผู้ดูแลระบบค่าเริ่มต้นเขียนไว้ในโค้ด", "api/views.py, admin_panel/views.py", "ย้ายไปไว้ใน .env และเปลี่ยนรหัสผ่านที่เคยเผยแพร่แล้ว"],
        ["มีฟิลด์เก็บรหัสผ่านแบบข้อความธรรมดา (password_plain)", "admin_panel/models.py, user_accounts.json", "ลบฟิลด์นี้และเก็บเฉพาะค่าแฮช"],
        ["แฮชรหัสผ่านด้วย SHA-256 แบบไม่มี salt", "views.py, api.php", "ใช้ bcrypt / Argon2 (password_hash() ของ PHP หรือ make_password ของ Django)"],
        ["ล็อกอินนักศึกษาฝั่ง PHP ยอมให้ผ่านถ้ารหัสนักศึกษาตรงกับ Roster แม้ยังไม่ลงทะเบียน", "auth/api.php (student_login)", "ตรวจรหัสผ่านทุกครั้ง หรือบังคับให้ลงทะเบียนก่อน"],
        ["เปิด CORS ให้ทุกโดเมน (Access-Control-Allow-Origin: *)", "auth/config.php", "จำกัดเฉพาะโดเมนของระบบ"],
        ["มี 2 โครงสร้างฐานข้อมูล (MySQL และ Django models) ที่ชื่อฟิลด์ไม่ตรงกัน", "migrate_schema.py, models.py", "เลือกโครงสร้างหลักเพียงชุดเดียว"],
    ], [2.3, 1.7, 2.2])
    doc.add_page_break()

    # ═══ 10. Traceability ═══
    h1(doc, "10. ตารางตรวจสอบย้อนกลับ (Traceability Matrix)")
    para(doc, "สถานะ 'พัฒนาแล้ว' หมายถึงพบโค้ดที่ทำงานตามความต้องการในโปรเจกต์ (ยังไม่ใช่ผลการทดสอบอย่างเป็นทางการ)", 11)
    table(doc, ["กลุ่มความต้องการ", "รหัส", "ไฟล์ที่เกี่ยวข้อง", "สถานะ"], [
        ["แผนที่และการแสดงผล", "FR-01 – FR-05", "app.js, index.php, styles.css", "พัฒนาแล้ว"],
        ["ค้นหาและรายละเอียดอาคาร", "FR-06 – FR-09", "app.js, buildings.json", "พัฒนาแล้ว"],
        ["GPS และนำทาง", "FR-10 – FR-15", "app.js", "พัฒนาแล้ว"],
        ["ยืนยันตัวตน + Guest Mode", "FR-16 – FR-23", "login.php, api.php, views.py, urls.py", "พัฒนาแล้ว"],
        ["โปรไฟล์และ PDPA", "FR-24 – FR-28", "app.js, index.php, pdpa_policy.json", "พัฒนาแล้ว"],
        ["ผู้ดูแลระบบ", "FR-29 – FR-35", "app.js, views.py, dashboard.html", "พัฒนาแล้ว"],
        ["ส่วนติดต่อผู้ใช้", "FR-36 – FR-39", "styles.css, dialog.js, manifest.json, sw.js", "พัฒนาแล้ว"],
        ["ความปลอดภัย", "NFR-04 – NFR-07", "views.py, api.php, .htaccess", "พัฒนาแล้ว (มีข้อเสนอแนะในหัวข้อ 9.2)"],
    ], [1.8, 1.2, 2.2, 1.0])

    out = os.path.join(BASE_DIR, "docs", "SSKRU_Campus_Map_Requirements.docx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print("SAVED:", out)


if __name__ == "__main__":
    build()
